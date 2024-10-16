import os
import uuid
import subprocess
from docx import Document
from django.conf import settings
from datetime import datetime

def manipular_docx(nome_aluno, nome_diretor, nome_escola, endereco_escola, data, media):
    # Caminho do template fixo do documento DOCX
    template_path = os.path.join(settings.BASE_DIR, 'contract', 'templates_documents', 'Carta.docx')

    # Verifique se o arquivo existe
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Arquivo {template_path} não encontrado.")

    # Caminho para a pasta media/documents/
    documents_dir = os.path.join(settings.MEDIA_ROOT, 'documents')

    # Criar o diretório se ele não existir
    if not os.path.exists(documents_dir):
        os.makedirs(documents_dir)

    # Carregar o template do documento
    doc = Document(template_path)

    # Formatar a data para o formato DD/MM/YYYY
    try:
        # Tentar converter a data no formato 'YYYY-MM-DD' para 'DD/MM/YYYY'
        data_formatada = datetime.strptime(data, '%Y-%m-%d').strftime('%d/%m/%Y')
    except ValueError:
        try:
            # Caso já esteja no formato 'DD/MM/YYYY', deixar como está
            data_formatada = datetime.strptime(data, '%d/%m/%Y').strftime('%d/%m/%Y')
        except ValueError:
            # Caso não consiga converter, usar a data original
            data_formatada = data

    # Formatar a nota para duas casas decimais
    media_formatada = f"{float(media):.2f}"

    # Substituir os placeholders no documento
    for p in doc.paragraphs:
        if '[Nome do Aluno]' in p.text:
            p.text = p.text.replace('[Nome do Aluno]', nome_aluno)
        if '[Nome do Diretor ou Diretora]' in p.text:
            p.text = p.text.replace('[Nome do Diretor ou Diretora]', nome_diretor)
        if '[Nome da Escola]' in p.text:
            p.text = p.text.replace('[Nome da Escola]', nome_escola)
        if '[Endereço Completo da Escola]' in p.text:
            p.text = p.text.replace('[Endereço Completo da Escola]', endereco_escola)
        if '[Data]' in p.text:
            p.text = p.text.replace('[Data]', data_formatada)
        if '[Nota]' in p.text:
            p.text = p.text.replace('[Nota]', media_formatada)

    # Gerar um UUID para o nome do arquivo
    unique_filename = f'{uuid.uuid4()}.docx'
    pdf_filename = unique_filename.replace('.docx', '.pdf')

    # Caminho para salvar o arquivo na pasta media/documents/
    docx_output_path = os.path.join(documents_dir, unique_filename)
    pdf_output_path = os.path.join(documents_dir, pdf_filename)

    # Salvar o documento manipulado como DOCX
    doc.save(docx_output_path)

    # Converter DOCX para PDF usando LibreOffice
    command = ['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', documents_dir, docx_output_path]
    subprocess.run(command, check=True)

    return pdf_output_path
